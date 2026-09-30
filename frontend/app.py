import os

import streamlit as st
import requests
import textwrap
from io import BytesIO
from docx import Document
from fpdf import FPDF

API_BASE_URL = os.getenv("API_BASE_URL", "http://127.0.0.1:8001").rstrip("/")
if not API_BASE_URL.startswith(("http://", "https://")):
    API_BASE_URL = f"https://{API_BASE_URL}"


# --------------------------------------------------
# Page Configuration
# --------------------------------------------------

st.set_page_config(
    page_title="LegalEase",
    page_icon="⚖️",
    layout="wide"
)

st.markdown(
    """
    <style>
        :root {
            --bg: #0b1220;
            --panel: #111b2d;
            --panel-strong: #16233d;
            --primary: #7dd3fc;
            --primary-strong: #38bdf8;
            --text: #e5edf9;
            --muted: #a5b4cf;
            --border: rgba(148, 163, 184, 0.2);
            --success: #34d399;
            --warning: #fbbf24;
        }

        html, body, [data-testid="stAppViewContainer"] {
            background: linear-gradient(135deg, #07111d 0%, #0d1728 35%, #101c30 100%);
            color: var(--text);
        }

        [data-testid="stHeader"] {
            background: rgba(11, 18, 32, 0.7);
            backdrop-filter: blur(10px);
        }

        .block-container {
            padding-top: 2rem;
            padding-bottom: 3rem;
        }

        .legal-header {
            background: linear-gradient(135deg, rgba(12, 74, 110, 0.82), rgba(59, 130, 246, 0.38), rgba(15, 23, 42, 0.7));
            border: 1px solid rgba(148, 163, 184, 0.22);
            border-radius: 24px;
            padding: 1.8rem 1.8rem;
            box-shadow: 0 20px 50px rgba(37, 99, 235, 0.18), inset 0 1px 0 rgba(255,255,255,0.08);
            margin-bottom: 1.4rem;
            backdrop-filter: blur(14px);
            -webkit-backdrop-filter: blur(14px);
        }

        .legal-header h1 {
            margin: 0;
            color: #f8fbff;
            font-size: 2.5rem;
            letter-spacing: 0.02em;
            text-shadow: 0 0 30px rgba(125, 211, 252, 0.45);
        }

        .legal-header p {
            margin: 0.5rem 0 0;
            color: rgba(226, 232, 240, 0.9);
            font-size: 1rem;
        }

        .info-card {
            background: rgba(15, 23, 42, 0.42);
            border: 1px solid rgba(148, 163, 184, 0.18);
            border-radius: 22px;
            padding: 1.2rem 1.3rem;
            box-shadow: 0 12px 30px rgba(2, 6, 23, 0.24), inset 0 1px 0 rgba(255,255,255,0.04);
            height: 100%;
            backdrop-filter: blur(10px);
            -webkit-backdrop-filter: blur(10px);
        }

        .floating-panel {
            background: rgba(15, 23, 42, 0.55);
            border: 1px solid rgba(148, 163, 184, 0.18);
            border-radius: 22px;
            padding: 1.2rem 1.3rem;
            box-shadow: 0 18px 38px rgba(15, 23, 42, 0.35);
            backdrop-filter: blur(12px);
            -webkit-backdrop-filter: blur(12px);
            margin-top: 1rem;
        }

        .sidebar {
            background: rgba(11, 18, 32, 0.7);
            border: 1px solid rgba(148, 163, 184, 0.16);
            border-radius: 22px;
            padding: 1rem 0.8rem;
            box-shadow: 0 10px 26px rgba(2, 6, 23, 0.18);
            position: sticky;
            top: 1rem;
        }

        .nav-item {
            display: block;
            color: rgba(226, 232, 240, 0.9);
            text-decoration: none;
            padding: 0.8rem 0.9rem;
            border-radius: 12px;
            margin: 0.35rem 0;
            background: transparent;
            border: 1px solid transparent;
            transition: all 0.2s ease;
        }

        .nav-item:hover {
            background: rgba(59, 130, 246, 0.08);
            border-color: rgba(125, 211, 252, 0.15);
        }

        .nav-item.active {
            background: linear-gradient(135deg, rgba(56, 189, 248, 0.18), rgba(59, 130, 246, 0.08));
            border-color: rgba(125, 211, 252, 0.3);
            color: #f8fbff;
        }

        .info-card h4 {
            margin-top: 0;
            color: var(--primary);
        }

        .info-card p, .info-card li {
            color: var(--muted);
            line-height: 1.6;
        }

        div[data-baseweb="select"] > div,
        textarea, input {
            background: rgba(15, 23, 42, 0.38) !important;
            color: var(--text) !important;
            border: 1px solid rgba(148, 163, 184, 0.18) !important;
            border-radius: 14px !important;
            box-shadow: inset 0 1px 0 rgba(255,255,255,0.02);
            backdrop-filter: blur(6px);
            -webkit-backdrop-filter: blur(6px);
        }

        .stTextInput label, .stTextArea label, .stSelectbox label {
            color: var(--text) !important;
            font-weight: 600;
        }

        .stButton > button {
            background: linear-gradient(135deg, var(--primary-strong), var(--primary));
            color: #06131f;
            border: none;
            border-radius: 12px;
            padding: 0.75rem 1.5rem;
            font-weight: 700;
            box-shadow: 0 10px 20px rgba(56, 189, 248, 0.25);
        }

        .stButton > button:hover {
            background: linear-gradient(135deg, #8ae2ff, #67d4ff);
            color: #051720;
        }

        .stAlert {
            background: rgba(15, 23, 42, 0.8);
            border: 1px solid var(--border);
            color: var(--text);
        }

        .stDownloadButton > button {
            background: rgba(17, 27, 45, 0.95);
            color: var(--text);
            border: 1px solid var(--border);
            border-radius: 12px;
        }
    </style>
    """,
    unsafe_allow_html=True,
)


# --------------------------------------------------
# Sidebar + Header
# --------------------------------------------------

sidebar_col, main_col = st.columns([0.22, 1.0])

with sidebar_col:
    st.markdown(
        """
        <div class="sidebar">
            <div class="nav-item active">🏠 Overview</div>
            <div class="nav-item">📄 Document Studio</div>
            <div class="nav-item">🧾 Templates</div>
            <div class="nav-item">📦 Exports</div>
            <div class="nav-item">⚙️ Settings</div>
        </div>
        """,
        unsafe_allow_html=True,
    )

with main_col:
    st.markdown(
        """
        <div class="legal-header">
            <h1>⚖️ LegalEase</h1>
            <p>AI-powered legal document generation for modern legal workflows.</p>
        </div>
        """,
        unsafe_allow_html=True,
    )

    left_col, right_col = st.columns([1.8, 0.9])

    with left_col:
        # --------------------------------------------------
        # Document Type
        # --------------------------------------------------

        document_type = st.selectbox(
            "Document Type",
            [
                "Non-Disclosure Agreement",
                "Employment Contract",
                "Freelance Work Contract",
                "Residential Lease Agreement",
                "Other"
            ]
        )


        # --------------------------------------------------
        # Parties
        # --------------------------------------------------

        parties = st.text_area(
            "Parties",
            placeholder="Example: Bhuvan (Freelancer), ABC Company (Client)",
            height=120
        )


        # --------------------------------------------------
        # Terms and Conditions
        # --------------------------------------------------

        terms = st.text_area(
            "Terms and Conditions",
            placeholder="Enter the important terms and conditions...",
            height=180
        )


        # --------------------------------------------------
        # Date
        # --------------------------------------------------

        dates = st.text_input(
            "Date",
            placeholder="Example: September 23, 2026"
        )

    with right_col:
        st.markdown(
            """
            <div class="info-card">
                <h4>Professional drafting workflow</h4>
                <p>Build contracts and agreements with clear structure, clause summaries, and export-ready outputs.</p>
                <ul>
                    <li>Fast legal drafting</li>
                    <li>Editable generated text</li>
                    <li>TXT, DOCX, and PDF exports</li>
                </ul>
            </div>
            """,
            unsafe_allow_html=True,
        )

    st.markdown(
        """
        <div class="floating-panel">
            <h4 style="margin-top:0; color: #f8fbff;">Document Builder</h4>
        </div>
        """,
        unsafe_allow_html=True,
    )


# --------------------------------------------------
# Generate Document
# --------------------------------------------------

if st.button("🚀 Generate Legal Document", type="primary", use_container_width=True):

    if not parties or not terms or not dates:

        st.warning("Please fill in all the fields.")

    else:

        with st.spinner("Generating your legal document..."):

            try:

                response = requests.post(
                    f"{API_BASE_URL}/generate",
                    json={
                        "document_type": document_type,
                        "parties": parties,
                        "terms": terms,
                        "dates": dates
                    }
                )


                # --------------------------------------------------
                # Successful Response
                # --------------------------------------------------

                if response.status_code == 200:

                    result = response.json()

                    document = result["document"]

                    st.success(
                        "Document generated successfully! 🎉"
                    )

                    st.subheader("Generated Document")


                    # --------------------------------------------------
                    # Preview / Edit Document
                    # --------------------------------------------------

                    edited_document = st.text_area(
                        "Preview / Edit Document",
                        value=document,
                        height=500
                    )


                    # --------------------------------------------------
                    # Download TXT
                    # --------------------------------------------------

                    st.download_button(
                        label="📄 Download TXT",
                        data=edited_document,
                        file_name="legal_document.txt",
                        mime="text/plain",
                        on_click="ignore"
                    )


                    # --------------------------------------------------
                    # Create DOCX
                    # --------------------------------------------------

                    doc = Document()

                    for paragraph in edited_document.split("\n"):
                        doc.add_paragraph(paragraph)

                    docx_file = BytesIO()

                    doc.save(docx_file)

                    docx_file.seek(0)


                    # --------------------------------------------------
                    # Download DOCX
                    # --------------------------------------------------

                    st.download_button(
                        label="📝 Download DOCX",
                        data=docx_file,
                        file_name="legal_document.docx",
                        mime=(
                            "application/vnd.openxmlformats-"
                            "officedocument.wordprocessingml.document"
                        ),
                        on_click="ignore"
                    )


                    # --------------------------------------------------
                    # Create PDF
                    # --------------------------------------------------

                    pdf = FPDF()

                    pdf.set_auto_page_break(
                        auto=True,
                        margin=15
                    )

                    pdf.add_page()


                    # --------------------------------------------------
                    # PDF Title
                    # --------------------------------------------------

                    pdf.set_font(
                        "Helvetica",
                        "B",
                        16
                    )

                    pdf.cell(
                        0,
                        10,
                        "LegalEase - Legal Document",
                        ln=1,
                        align="C"
                    )

                    pdf.ln(5)


                    # --------------------------------------------------
                    # PDF Body
                    # --------------------------------------------------

                    pdf.set_font(
                        "Helvetica",
                        "",
                        11
                    )

                    for line in edited_document.split("\n"):

                        # Remove simple Markdown formatting
                        line = line.replace("**", "")
                        line = line.replace("### ", "")
                        line = line.replace("## ", "")
                        line = line.replace("# ", "")

                        # Convert unsupported characters
                        line = line.encode(
                            "latin-1",
                            "replace"
                        ).decode("latin-1")

                        if line.strip():
                            wrapped_lines = textwrap.wrap(
                                line,
                                width=95,
                                break_long_words=True,
                                break_on_hyphens=False,
                                replace_whitespace=False
                            )

                            if not wrapped_lines:
                                wrapped_lines = [line[:95]]

                            for wrapped_line in wrapped_lines:
                                pdf.multi_cell(
                                    pdf.w - 2 * pdf.l_margin,
                                    7,
                                    wrapped_line
                                )

                        else:

                            pdf.ln(4)


                    # --------------------------------------------------
                    # Convert PDF to bytes
                    # --------------------------------------------------

                    pdf_output = pdf.output(dest='S')

                    if isinstance(pdf_output,str):
                        pdf_bytes = pdf_output.encode("latin-1")
                    else:
                        pdf_bytes = bytes(pdf_output)

                    # --------------------------------------------------
                    # Download PDF
                    # --------------------------------------------------

                    st.download_button(
                        label="📕 Download PDF",
                        data=pdf_bytes,
                        file_name="legal_document.pdf",
                        mime="application/pdf",
                        on_click="ignore"
                    )


                # --------------------------------------------------
                # Backend Error
                # --------------------------------------------------

                else:
                    try:
                        error_detail = response.json().get("detail", response.text)
                    except Exception:
                        error_detail = response.text or f"Backend error: {response.status_code}"

                    st.error(
                        f"Backend error: {response.status_code}\n\n{error_detail}"
                    )


            # --------------------------------------------------
            # Connection Error
            # --------------------------------------------------

            except requests.exceptions.ConnectionError:

                st.error(
                    "Could not connect to the FastAPI backend. "
                    "Make sure the backend is running."
                )